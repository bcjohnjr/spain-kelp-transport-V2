from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
out=Path('/mnt/data/Spain_Nature_Sustainability_Dual_Path_V16')
main=out/'Submission'/'Manuscript_Nature_Sustainability_Dual_Path_V16.docx'
d=Document(main)
# fix citation 29 with verified author metadata
for p in d.paragraphs:
 if p.text.startswith('29. Aging processes and microplastic release'):
  p.text='29. Duan, Y. et al. Aging processes and microplastic release behavior of aquaculture implements. Environ. Pollut. 392, 127620 (2026). https://doi.org/10.1016/j.envpol.2025.127620'
 if 'Corresponding-author institutional details and coauthor approvals require author verification' in p.text:
  p.text=p.text.replace(' Corresponding-author institutional details and coauthor approvals require author verification before submission.','')
# supporting text for falsifiable novelty: 2 new paragraphs, not generative art
for anchor,text in [
 ('The proposed autonomous *selective-hook* drone is a central',
 'The engineering protocol should randomize comparable cultivation sections between conventional capture and robotic selective capture, measuring fish caught per device-hour, size-selectivity error, escaped or injured fish, post-release survival, harvest interruption time, kelp damage, maintenance burden and energy per kilogram landed. Autonomous operation in fouled, low-visibility pens may be particularly challenging. The national fish-margin figure assumes pre-existing harvest economics and therefore does not demonstrate a robotic advantage; separate equipment depreciation, charging, failure and animal-welfare costs would need to be deducted after trials. A purpose-built hook may be technically distinguishable, but no novelty or patentability claim follows from the present model.'),
 ('Coastal collection of abandoned fishing gear and macroplastic',
 'A materials balance should begin with kilograms of recovered macroplastic, identify actual polymer grade and usable yield after washing and contamination rejection, and separately estimate fabrication energy, steel-core cost and recoverable float volume. A structural prototype should be load-tested through repeated flexure, salinity exposure and storm-force cycles, with abrasion-related microplastics collected and quantified. The economic comparator must include purchase or leasing of ordinary marine floats, maintenance, loss and retrieval obligations; potential avoided virgin polymer does not itself prove a cost reduction. Preferring detachable, inspectable components also permits replacement when a jacket fails, without discarding a serviceable steel core. These metrics are necessary to establish whether the circular materials proposal improves on ordinary shore-based plastics recycling.')]:
 p=next(p for p in d.paragraphs if p.text.startswith(anchor))
 n=OxmlElement('w:p');p._p.addnext(n);q=Paragraph(n,p._parent);q.text=text;q.style=d.styles['Normal']
# count word limit
end=next(i for i,p in enumerate(d.paragraphs) if p.text=='Methods')
body=[p for p in d.paragraphs[6:end] if p.text and p.text not in ('Results','Discussion') and not p.style.name.startswith('Heading') and not p.text.startswith(('Fig. ', 'Table 1 |','Notes:'))]
wcount=sum(len(p.text.split()) for p in body)
print('MAIN word count:',wcount,'abstract:',len(d.paragraphs[5].text.split()),'figs:',len(d.inline_shapes),'tables:',len(d.tables))
if wcount>3500:
 # trim surplus by shortening a non-critical paragraph, without hiding results
 p=next(p for p in d.paragraphs if p.text.startswith('This separation does not imply that'))
 p.text='Consumer vehicle purchases remain outside the energy-supplier boundary; this does not mean that electrification carries no private cost.'
 end=next(i for i,p in enumerate(d.paragraphs) if p.text=='Methods')
 body=[p for p in d.paragraphs[6:end] if p.text and p.text not in ('Results','Discussion') and not p.style.name.startswith('Heading') and not p.text.startswith(('Fig. ', 'Table 1 |','Notes:'))]
 wcount=sum(len(p.text.split()) for p in body)
 print('trimmed',wcount)
assert wcount<=3500
# keep known wordcount note
(out/'WORD_COUNT_V16.txt').write_text(f'Main-text paragraph word count (Intro, Results, Discussion; excluding Methods, abstract, figure captions, Table and references): {wcount}.\nAbstract: {len(d.paragraphs[5].text.split())}.\nDisplay items: {len(d.inline_shapes)} figures + {len(d.tables)} table = {len(d.inline_shapes)+len(d.tables)}.\n')
d.save(main)
# reference special case in supplement
si=out/'Supplementary'/'Supplementary_Information_Dual_Path_V16.docx'
s=Document(si)
for p in s.paragraphs:
 if 'The earlier 2,110 figure' in p.text:print('check old drone count',p.text[:80])
s.save(si)
