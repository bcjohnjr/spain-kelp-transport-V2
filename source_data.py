"""Write one Source Data workbook per main-text figure (outputs/source_data/Source_Data_FigN.xlsx).

Each workbook has one sheet per panel holding the exact values plotted, as recorded by the figure scripts.
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

D = Path(__file__).resolve().parent / "outputs" / "source_data"


def main():
    for f in sorted(D.glob("Fig*.json")):
        data = json.load(open(f))
        wb = Workbook(); wb.remove(wb.active)
        for panel, tab in data.items():
            ws = wb.create_sheet(panel[:31])
            r0 = 1
            if tab.get("note"):
                ws.cell(1, 1, tab["note"]).font = Font(italic=True); r0 = 3
            for j, c in enumerate(tab["columns"], 1):
                ws.cell(r0, j, c).font = Font(bold=True)
            for i, row in enumerate(tab["rows"], r0 + 1):
                for j, v in enumerate(row, 1):
                    ws.cell(i, j, v)
            for j, c in enumerate(tab["columns"], 1):
                ws.column_dimensions[ws.cell(r0, j).column_letter].width = max(12, min(48, len(str(c)) + 2))
        out = D / f"Source_Data_{f.stem}.xlsx"
        wb.save(out); print("wrote", out.name)


if __name__ == "__main__":
    main()
