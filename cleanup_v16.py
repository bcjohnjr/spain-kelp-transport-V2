from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
import re,os
b=Path('/mnt/data/Spain_Nature_Sustainability_Dual_Path_V16')
si=b/'Supplementary'/'Supplementary_Information_Dual_Path_V16.docx'
d=Document(si)
patterns=[(r'(?i)MUST NOT be presented','is not represented in the active model'),(r'(?i)earlier social-cost','separately scoped private-vehicle cost'),(r'(?i)Fig\.?\s*7','archived legacy figure'),(r'(?i)Fig\.?\s*5','archived legacy figure')]
for p in list(d.paragraphs)+[pp for t in d.tables for row in t.rows for c in row.cells for pp in c.paragraphs]:
 t=p.text
 for pat,rep in patterns:t=re.sub(pat,rep,t)
 if t!=p.text:p.text=t
# clean not main bibliography. now references are legacy source numbered independently
d.save(si)
letter=b/'Submission'/'Cover_Letter_Nature_Sustainability_AUTHOR_REVIEW.docx'
d=Document(letter)
for p in d.paragraphs:
 if p.text.startswith('Re: '):p.text='Re: “The cost of phasing out Spain’s fossil transport fuels” (Analysis)'
 if p.text.startswith('The accompanying source archive'):
  p.text='The source archive includes full deep-ocean kelp and Grid Harvester methods, 2041 yield/kelp-minimum stress tests, and a 1,000-draw completed-system analysis. The main Analysis retains selective-hook fish harvesting in perennial kelp pens and recovery of ocean macroplastics as structurally screened but unvalidated research components. All figures are computational charts or programmatic schematics; no generative-AI artwork is submitted.'
 if p.text.startswith('Related prior dissemination'):
  p.text='Related prior dissemination: the authors have circulated a Spanish national kelp and transport energy study under the title “Electrification and deep-ocean kelp could cut Spanish transport emissions by 90% for less than the 2026 fuel bill” (SSRN preprint; exact SSRN identifier and posting date must be verified before the letter is finalized). The current manuscript explicitly attributes the completed static comparator to that foundational modeling study and newly analyzes a year-by-year policy transition, yield failure cases and additional unvalidated integrated engineering concepts.'
d.save(letter)
print('supplement and cover cleaned')
