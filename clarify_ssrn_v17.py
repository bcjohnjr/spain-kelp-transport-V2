"""Final SSRN prior-work clarification for the generated V17 letter.
Run AFTER build_v17.py and BEFORE rendering PDFs. Preserves the narrower-preprint disclosure.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
base=Path('/mnt/data/Spain_Nature_Sustainability_Dual_Path_V17')
letter=base/'Submission/Cover_Letter_Nature_Sustainability_V17.docx'
d=Document(letter)
for p in d.paragraphs:
 if p.text.startswith(('Related research has been disseminated','Prior dissemination')):
  p.text=('Prior dissemination: An earlier, narrower article, “Electrification and deep-ocean kelp could cut Spanish transport emissions by 90% for less than the 2026 fuel bill,” was posted as an SSRN preprint. That study modeled a completed, single-year electrification–kelp system. It did not include the present 2027–2041 fleet and infrastructure transition, annual and cumulative transition economics, supplier-cost reconciliation, or kelp-yield failure scenarios. It also lacked the present proposed wheat-fed hatchery fish co-culture, selective-hook underwater harvesting, and recovered-ocean-plastics flotation and protection concepts. We explicitly attribute the reused static optimizer and its Monte Carlo draws to that earlier work; the additional analyses and unvalidated engineering proposals are distinguished from it. The exact SSRN identifier and posting date remain unverified.')
 if p.text.startswith('The source archive includes full deep-ocean kelp'):
  p.text=('We supply reproducible model code, yield-stress tests and detailed kelp and Grid Harvester methods. The main text retains the proposed fishing-drone and marine-plastics pathways, clearly labeled as unvalidated. All submitted figures are data plots or programmatic schematics, not AI-generated artwork.')
for sec in d.sections:
 sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.66)
 sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
for i,p in enumerate(d.paragraphs):
 f=p.paragraph_format
 f.space_before=Pt(0);f.space_after=Pt(5 if i else 11); f.line_spacing=1.05
 for r in p.runs:r.font.size=Pt(16 if i==0 else 10.5)
d.save(letter)
